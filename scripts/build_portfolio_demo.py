"""Build the CRM Customer Database portfolio demonstration workbook."""

from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parent.parent

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


from openpyxl import load_workbook

from builder.activity_log import ACTIVITY_LOG_TABLE_NAME
from builder.constants import SheetNames
from builder.customer_database import CUSTOMER_TABLE_NAME
from builder.workbook_builder import build_workbook


SOURCE_DEMO = (
    PROJECT_ROOT
    / "output"
    / "Customer_Data_Management_Portfolio_Demo.xlsx"
)

OUTPUT_DEMO = (
    PROJECT_ROOT
    / "output"
    / "Customer_Data_Management_Portfolio_Demo.xlsx"
)

TEMP_SOURCE = (
    PROJECT_ROOT
    / "output"
    / "_portfolio_demo_source.xlsx"
)

CUSTOMER_FIRST_DATA_ROW = 5
CUSTOMER_LAST_DATA_ROW = 104
CUSTOMER_COLUMNS = 10
EXPECTED_RECORD_COUNT = 100

ACTIVITY_FIRST_DATA_ROW = 5
ACTIVITY_LAST_DATA_ROW = 9
EXPECTED_ACTIVITY_COUNT = 5

DEMO_ACTIVITIES: tuple[tuple[object, ...], ...] = (
    (
        "ACT-0001",
        datetime(2026, 8, 20, 9, 15),
        "CUST-0001",
        "New Customer",
        "Customer record created and initial information entered.",
        "Portfolio Demo",
    ),
    (
        "ACT-0002",
        datetime(2026, 8, 21, 10, 30),
        "CUST-0025",
        "Data Review",
        "Customer information reviewed for completeness and accuracy.",
        "Portfolio Demo",
    ),
    (
        "ACT-0003",
        datetime(2026, 8, 22, 13, 45),
        "CUST-0050",
        "Contact Update",
        "Customer contact information reviewed and updated.",
        "Portfolio Demo",
    ),
    (
        "ACT-0004",
        datetime(2026, 8, 24, 14, 20),
        "CUST-0075",
        "Follow-up",
        "Follow-up activity recorded for customer communication.",
        "Portfolio Demo",
    ),
    (
        "ACT-0005",
        datetime(2026, 8, 25, 16, 10),
        "CUST-0100",
        "Data Quality Check",
        "Customer record checked as part of the data-quality review.",
        "Portfolio Demo",
    ),
)


def _read_customer_records(
    source_path: Path,
) -> list[list[object]]:
    """Read verified customer records from the source workbook."""
    if not source_path.exists():
        raise FileNotFoundError(
            f"Source portfolio demo does not exist: {source_path}"
        )

    workbook = load_workbook(
        source_path,
        data_only=False,
    )

    try:
        worksheet = workbook[SheetNames.customer_database]

        if CUSTOMER_TABLE_NAME not in worksheet.tables:
            raise ValueError(
                f"Source workbook does not contain "
                f"{CUSTOMER_TABLE_NAME}."
            )

        table = worksheet.tables[CUSTOMER_TABLE_NAME]

        if table.ref != "A4:J104":
            raise ValueError(
                f"Unexpected source customer table range: "
                f"{table.ref}. Expected A4:J104."
            )

        records = [
            list(row)
            for row in worksheet.iter_rows(
                min_row=CUSTOMER_FIRST_DATA_ROW,
                max_row=CUSTOMER_LAST_DATA_ROW,
                min_col=1,
                max_col=CUSTOMER_COLUMNS,
                values_only=True,
            )
        ]

        if len(records) != EXPECTED_RECORD_COUNT:
            raise ValueError(
                f"Expected {EXPECTED_RECORD_COUNT} customer records, "
                f"found {len(records)}."
            )

        return records

    finally:
        workbook.close()


def _validate_source(source_path: Path) -> None:
    """Validate the existing portfolio source workbook."""
    records = _read_customer_records(source_path)

    if records[0][0] != "CUST-0001":
        raise ValueError(
            f"Unexpected first customer ID: {records[0][0]!r}."
        )

    if records[-1][0] != "CUST-0100":
        raise ValueError(
            f"Unexpected last customer ID: {records[-1][0]!r}."
        )


def _build_demo(
    records: list[list[object]],
    output_path: Path,
) -> Path:
    """Build the production workbook and populate demo records."""
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if output_path.exists():
        output_path.unlink()

    build_workbook(output_path)

    workbook = load_workbook(
        output_path,
        data_only=False,
    )

    try:
        customer_sheet = workbook[SheetNames.customer_database]
        activity_sheet = workbook[SheetNames.activity_log]

        for row_offset, record in enumerate(records):
            row_number = CUSTOMER_FIRST_DATA_ROW + row_offset

            for column_number, value in enumerate(
                record,
                start=1,
            ):
                customer_sheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value,
                )

        customer_table = customer_sheet.tables[
            CUSTOMER_TABLE_NAME
        ]
        customer_table.ref = (
            f"A4:J{CUSTOMER_LAST_DATA_ROW}"
        )

        for row_offset, activity in enumerate(
            DEMO_ACTIVITIES
        ):
            row_number = ACTIVITY_FIRST_DATA_ROW + row_offset

            for column_number, value in enumerate(
                activity,
                start=1,
            ):
                activity_sheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value,
                )

        activity_table = activity_sheet.tables[
            ACTIVITY_LOG_TABLE_NAME
        ]
        activity_table.ref = (
            f"A4:F{ACTIVITY_LAST_DATA_ROW}"
        )

        workbook.save(output_path)

    finally:
        workbook.close()

    return output_path


def _validate_output(output_path: Path) -> None:
    """Validate the completed portfolio demo workbook."""
    workbook = load_workbook(
        output_path,
        data_only=False,
    )

    try:
        expected_sheets = [
            SheetNames.cover,
            SheetNames.validation_lists,
            SheetNames.customer_database,
            SheetNames.dashboard,
            SheetNames.data_quality,
            SheetNames.activity_log,
            SheetNames.instructions,
        ]

        if workbook.sheetnames != expected_sheets:
            raise AssertionError(
                f"Unexpected sheets: {workbook.sheetnames}"
            )

        customer_sheet = workbook[
            SheetNames.customer_database
        ]
        activity_sheet = workbook[
            SheetNames.activity_log
        ]

        customer_table = customer_sheet.tables[
            CUSTOMER_TABLE_NAME
        ]
        activity_table = activity_sheet.tables[
            ACTIVITY_LOG_TABLE_NAME
        ]

        if customer_table.ref != "A4:J104":
            raise AssertionError(
                f"Unexpected customer table range: "
                f"{customer_table.ref}"
            )

        if activity_table.ref != "A4:F9":
            raise AssertionError(
                f"Unexpected Activity Log table range: "
                f"{activity_table.ref}"
            )

        activity_records = [
            [
                activity_sheet.cell(row, column).value
                for column in range(1, 7)
            ]
            for row in range(
                ACTIVITY_FIRST_DATA_ROW,
                ACTIVITY_LAST_DATA_ROW + 1,
            )
        ]

        if len(activity_records) != EXPECTED_ACTIVITY_COUNT:
            raise AssertionError(
                f"Expected {EXPECTED_ACTIVITY_COUNT} activity records, "
                f"found {len(activity_records)}."
            )

        if activity_records != [
            list(activity)
            for activity in DEMO_ACTIVITIES
        ]:
            raise AssertionError(
                "Portfolio demo Activity Log records do not "
                "match the expected demo activities."
            )

        validation_count = len(
            customer_sheet.data_validations.dataValidation
        )

        if validation_count != 5:
            raise AssertionError(
                "Expected 5 customer data validations."
            )

        if len(workbook.defined_names) != 5:
            raise AssertionError(
                "Expected 5 defined names."
            )

        if activity_sheet.freeze_panes != "A5":
            raise AssertionError(
                f"Unexpected Activity Log freeze panes: "
                f"{activity_sheet.freeze_panes}"
            )

        if customer_sheet.cell(
            CUSTOMER_FIRST_DATA_ROW,
            1,
        ).value != "CUST-0001":
            raise AssertionError(
                "First customer record is not CUST-0001."
            )

        if customer_sheet.cell(
            CUSTOMER_LAST_DATA_ROW,
            1,
        ).value != "CUST-0100":
            raise AssertionError(
                "Last customer record is not CUST-0100."
            )

        customer_ids = [
            customer_sheet.cell(row, 1).value
            for row in range(
                CUSTOMER_FIRST_DATA_ROW,
                CUSTOMER_LAST_DATA_ROW + 1,
            )
        ]

        if len(customer_ids) != EXPECTED_RECORD_COUNT:
            raise AssertionError(
                f"Expected {EXPECTED_RECORD_COUNT} customer IDs, "
                f"found {len(customer_ids)}."
            )

        if len(set(customer_ids)) != EXPECTED_RECORD_COUNT:
            raise AssertionError(
                "Customer IDs are not unique."
            )

        if customer_sheet.max_row != CUSTOMER_LAST_DATA_ROW:
            raise AssertionError(
                f"Unexpected customer worksheet max row: "
                f"{customer_sheet.max_row}"
            )

    finally:
        workbook.close()


def main() -> None:
    """Build and validate the portfolio demonstration workbook."""
    if not SOURCE_DEMO.exists():
        raise FileNotFoundError(
            f"Source portfolio demo not found: {SOURCE_DEMO}"
        )

    print(
        "1. Creating temporary copy of the existing demo..."
    )

    if TEMP_SOURCE.exists():
        TEMP_SOURCE.unlink()

    shutil.copy2(
        SOURCE_DEMO,
        TEMP_SOURCE,
    )

    try:
        print(
            "2. Validating the existing 100-record source..."
        )

        _validate_source(TEMP_SOURCE)

        print("   Source validation passed.")

        print("3. Reading customer records...")

        records = _read_customer_records(
            TEMP_SOURCE
        )

        print(
            f"   Records loaded: {len(records)}"
        )

        print(
            "4. Building current production workbook..."
        )

        _build_demo(
            records,
            OUTPUT_DEMO,
        )

        print(
            "5. Validating completed portfolio demo..."
        )

        _validate_output(OUTPUT_DEMO)

        print()
        print(
            "PORTFOLIO DEMO BUILD SUCCESSFUL"
        )
        print(
            f"Output: {OUTPUT_DEMO}"
        )
        print(
            "Customer records: 100"
        )
        print(
            "Customer table: A4:J104"
        )
        print(
            "Activity records: 5"
        )
        print(
            "Activity table: A4:F9"
        )
        print(
            "Customer validations: 5"
        )
        print(
            "Defined names: 5"
        )

    finally:
        if TEMP_SOURCE.exists():
            TEMP_SOURCE.unlink()


if __name__ == "__main__":
    main()